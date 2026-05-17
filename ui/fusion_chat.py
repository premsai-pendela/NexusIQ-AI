"""
NexusIQ AI — Fusion Agent Chat Interface
Multi-source intelligence combining SQL + RAG + Web with cross-validation

Features:
- Smart routing display (SQL/RAG/Web/Combined)
- Multi-source result sections (expandable)
- Cross-validation confidence badges
- Query history with full fusion results
- Multi-format export (CSV/JSON/Excel/MD)
- Chart builder for SQL data
- Source filters + category selector
- All existing SQL chat features preserved
"""

'''import streamlit as st
import pandas as pd
import time
import random
import sys
import json
import io
import threading
from streamlit.runtime.scriptrunner import add_script_run_ctx, get_script_run_ctx
from datetime import datetime
from pathlib import Path
import plotly.express as px
import plotly.graph_objects as go

sys.path.append(str(Path(__file__).parent.parent))

from observability.inspect_traces import SLOW_SPAN_SECONDS, get_trace_diagnostics

from agents.fusion_agent import get_fusion_agent
from config.settings import settings
from utils.validators import VALID_REGIONS, VALID_CATEGORIES'''

import streamlit as st
import time
import random
import sys
import json
import io
import threading
from datetime import datetime
from pathlib import Path

sys.path.append(str(Path(__file__).parent.parent))

from observability.inspect_traces import SLOW_SPAN_SECONDS, get_trace_diagnostics

# ═══════════════════════════════════════════════════════
#  CONFIGURATION
# ═══════════════════════════════════════════════════════

INSIGHTS = [
    "💡 **Did You Know?** Fusion Agent combines 3 data sources for validated answers",
    "🧠 **Did You Know?** Cross-validation checks SQL numbers against PDF reports",
    "⚡ **Did You Know?** Simple queries complete in under 5 seconds",
    "🔄 **Did You Know?** We auto-switch models if one hits quota limits",
    "🔒 **Did You Know?** All SQL queries are read-only — your data stays safe",
    "📊 **Did You Know?** RAG Agent searches 23 business documents",
    "🎯 **Did You Know?** Web Agent scrapes live competitor pricing",
    "🚀 **Did You Know?** Circuit breaker skips failed models instantly",
    "💰 **Did You Know?** Database tracks $139M+ in total revenue",
    "🌐 **Did You Know?** We support 5 regions: East, West, North, South, Central",
    "🛒 **Did You Know?** Web Agent supports 5 categories: Electronics, Home, Sports, Food, Clothing",
    "⏱️ **Did You Know?** First query may be slower due to model warm-up",
    "🔍 **Did You Know?** Fusion Agent detects comparison queries and uses multi-step reasoning",
    "📄 **Did You Know?** RAG Agent uses hybrid BM25+Vector search for better accuracy",
]

CHART_TYPES = {
    "bar": {"icon": "📊", "name": "Bar Chart", "description": "Compare categories"},
    "bar_horizontal": {"icon": "📊", "name": "Horizontal Bar", "description": "Ranking/Top N"},
    "line": {"icon": "📈", "name": "Line Chart", "description": "Trends over time"},
    "pie": {"icon": "🥧", "name": "Pie Chart", "description": "Show proportions"},
    "scatter": {"icon": "🔵", "name": "Scatter Plot", "description": "Find patterns"},
    "area": {"icon": "📉", "name": "Area Chart", "description": "Cumulative trends"},
}

SOURCE_ICONS = {
    "sql": "🗄️",
    "rag": "📄",
    "web": "🌐",
    "fusion": "🔗"
}

RECRUITER_DEMO_PROMPTS = [
    {
        "title": "Validate Q4 Electronics revenue",
        "question": "Validate Q4 Electronics revenue across SQL and PDF reports.",
        "route": "SQL + RAG",
        "proof": "Shows exact database value, document value, source difference, and confidence.",
    },
    {
        "title": "Compare Q3 and Q4 performance",
        "question": "Compare Q3 and Q4 2024 performance across all metrics.",
        "route": "RAG comparison",
        "proof": "Demonstrates hybrid retrieval and multi-document synthesis.",
    },
    {
        "title": "Explain West vs South",
        "question": "Explain why West region outperformed South in 2024.",
        "route": "SQL + RAG",
        "proof": "Combines region numbers with strategic and market context.",
    },
    {
        "title": "Find competitor pricing",
        "question": "What are competitor prices for electronics?",
        "route": "Live web",
        "proof": "Uses external market data instead of only prepared local documents.",
    },
    {
        "title": "Summarize return policy",
        "question": "What is the return policy for Electronics?",
        "route": "RAG",
        "proof": "Retrieves internal policy details with document citations.",
    },
    {
        "title": "Count October transactions",
        "question": "Show October 2024 transaction count from the database.",
        "route": "SQL",
        "proof": "Runs an exact structured query over transaction data.",
    },
]

# ═══════════════════════════════════════════════════════
#  LAZY LOADERS — heavy modules loaded only when needed
# ═══════════════════════════════════════════════════════

def _get_pd():
    """Lazy import pandas"""
    import pandas as pd
    return pd

def _get_px():
    """Lazy import plotly express"""
    import plotly.express as px
    return px

def _get_go():
    """Lazy import plotly graph objects"""
    import plotly.graph_objects as go
    return go

def _get_settings():
    """Lazy import settings"""
    from config.settings import settings
    return settings

def _get_validators():
    """Lazy import validators"""
    from utils.validators import VALID_REGIONS, VALID_CATEGORIES
    return VALID_REGIONS, VALID_CATEGORIES

# ═══════════════════════════════════════════════════════
#  HELPER FUNCTIONS (from sql_chat.py - kept as-is)
# ═══════════════════════════════════════════════════════

def format_time(seconds: float) -> str:
    if seconds < 60:
        return f"{int(seconds)} seconds"
    elif seconds < 3600:
        mins = int(seconds // 60)
        secs = int(seconds % 60)
        if secs == 0:
            return f"{mins} minute{'s' if mins > 1 else ''}"
        return f"{mins} minute{'s' if mins > 1 else ''} {secs} seconds"
    else:
        hours = int(seconds // 3600)
        remaining = seconds % 3600
        mins = int(remaining // 60)
        secs = int(remaining % 60)
        result = f"{hours} hour{'s' if hours > 1 else ''}"
        if mins > 0:
            result += f" {mins} minute{'s' if mins > 1 else ''}"
        if secs > 0:
            result += f" {secs} seconds"
        return result

def time_ago(timestamp: datetime) -> str:
    diff = datetime.now() - timestamp
    seconds = diff.total_seconds()
    if seconds < 60:
        return "just now"
    elif seconds < 3600:
        return f"{int(seconds // 60)}m ago"
    elif seconds < 86400:
        return f"{int(seconds // 3600)}h ago"
    else:
        return f"{int(seconds // 86400)}d ago"

def _humanize_column_name(col: str) -> str:
    return str(col).replace("_", " ").strip().title()

def _format_metric_value(col: str, value) -> str:
    if value is None:
        return "N/A"
    col_lower = str(col).lower()
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)

    if any(term in col_lower for term in ["revenue", "amount", "sales", "price", "cost"]):
        return f"${number:,.2f}"
    if any(term in col_lower for term in ["rate", "margin", "percent", "pct"]):
        return f"{number:,.2f}%"
    if number.is_integer():
        return f"{int(number):,}"
    return f"{number:,.2f}"

def prepare_visualization_dataframe(df):
    """Return a plotting copy with numeric-looking object columns coerced."""
    pd = _get_pd()
    plot_df = df.copy()
    numeric_cols = plot_df.select_dtypes(include=['number']).columns.tolist()

    for col in plot_df.columns:
        if col in numeric_cols:
            continue
        converted = pd.to_numeric(plot_df[col], errors="coerce")
        non_null = plot_df[col].notna().sum()
        if non_null and converted.notna().sum() == non_null:
            plot_df[col] = converted
            numeric_cols.append(col)

    return plot_df, numeric_cols

def render_kpi_summary(df) -> bool:
    """Render single-row aggregate results as metrics. Returns True if rendered."""
    if df is None or df.empty or len(df) != 1:
        return False

    plot_df, numeric_cols = prepare_visualization_dataframe(df)
    if not numeric_cols:
        return False

    st.markdown("**📌 Key Metrics:**")
    metric_cols = st.columns(min(len(numeric_cols), 4))
    row = plot_df.iloc[0]
    for idx, col in enumerate(numeric_cols):
        with metric_cols[idx % len(metric_cols)]:
            st.metric(_humanize_column_name(col), _format_metric_value(col, row[col]))

    return True

def can_visualize(df) -> dict:
    """Check if dataframe can be visualized"""
    if df is None or df.empty:
        return {
            "can_chart": False,
            "reason": "No data to visualize",
            "numeric_cols": [],
            "text_cols": [],
            "date_cols": []
        }

    if len(df) < 1:
        return {
            "can_chart": False,
            "reason": "Need at least 1 row of data",
            "numeric_cols": [],
            "text_cols": [],
            "date_cols": []
        }

    plot_df, numeric_cols = prepare_visualization_dataframe(df)
    text_cols = df.select_dtypes(include=['object', 'category']).columns.tolist()
    date_cols = [
        col for col in df.columns
        if 'date' in col.lower() or 'month' in col.lower() 
        or 'year' in col.lower() or 'time' in col.lower()
    ]

    if len(df) == 1:
        return {
            "can_chart": False,
            "reason": "Single-row aggregate shown as KPI",
            "numeric_cols": numeric_cols,
            "text_cols": text_cols,
            "date_cols": date_cols,
            "plot_df": plot_df
        }

    if not numeric_cols:
        return {
            "can_chart": False,
            "reason": "No chartable numeric fields found",
            "numeric_cols": [],
            "text_cols": text_cols,
            "date_cols": date_cols,
            "plot_df": plot_df
        }

    return {
        "can_chart": True,
        "reason": "Ready to visualize!",
        "numeric_cols": numeric_cols,
        "text_cols": text_cols,
        "date_cols": date_cols,
        "row_count": len(df),
        "plot_df": plot_df
    }

def generate_chart(df, chart_type: str, x_col: str, y_col: str, color_col: str = None):
    px = _get_px()
    go = _get_go()
    try:
        title = f"{y_col} by {x_col}"
        
        if chart_type == "bar":
            fig = px.bar(
                df, x=x_col, y=y_col,
                color=color_col if color_col and color_col != "None" else None,
                title=f"📊 {title}",
                text_auto=True
            )
            
        elif chart_type == "bar_horizontal":
            fig = px.bar(
                df, x=y_col, y=x_col,
                color=color_col if color_col and color_col != "None" else None,
                title=f"📊 {title}",
                orientation='h',
                text_auto=True
            )
            fig.update_layout(yaxis={'categoryorder': 'total ascending'})
            
        elif chart_type == "line":
            fig = px.line(
                df, x=x_col, y=y_col,
                color=color_col if color_col and color_col != "None" else None,
                title=f"📈 {title}",
                markers=True
            )
            
        elif chart_type == "pie":
            fig = px.pie(
                df, names=x_col, values=y_col,
                title=f"🥧 {title}",
                hole=0.4
            )
            
        elif chart_type == "scatter":
            fig = px.scatter(
                df, x=x_col, y=y_col,
                color=color_col if color_col and color_col != "None" else None,
                title=f"🔵 {title}",
                size=y_col if len(df) > 1 else None
            )
            
        elif chart_type == "area":
            fig = px.area(
                df, x=x_col, y=y_col,
                color=color_col if color_col and color_col != "None" else None,
                title=f"📉 {title}"
            )
        
        else:
            fig = px.bar(df, x=x_col, y=y_col, title=title)

        fig.update_layout(
            template="plotly_white",
            height=400,
            showlegend=bool(color_col and color_col != "None"),
            margin=dict(t=50, b=50, l=50, r=50)
        )
        
        return fig
        
    except Exception as e:
        fig = go.Figure()
        fig.add_annotation(
            text=f"⚠️ Chart Error: {str(e)}",
            xref="paper", yref="paper",
            x=0.5, y=0.5, showarrow=False,
            font=dict(size=16, color="red")
        )
        fig.update_layout(height=200)
        return fig

# ─────────────────────────────────────────────────────
#  CHART BUILDER UI (from sql_chat.py - kept as-is)
# ─────────────────────────────────────────────────────

def render_chart_builder(msg_id: str, df):
    """Render the chart builder interface for a message."""
    
    viz_info = can_visualize(df)
    
    if not viz_info["can_chart"]:
        if viz_info["reason"] != "Single-row aggregate shown as KPI":
            st.caption(f"📊 {viz_info['reason']}")
        return None
    
    df = viz_info["plot_df"]

    st.markdown("**🎨 Build Your Chart**")
    
    chart_cols = st.columns(6)
    selected_chart = st.session_state.get(f"chart_type_{msg_id}", "bar")
    
    for i, (chart_key, chart_info) in enumerate(CHART_TYPES.items()):
        with chart_cols[i]:
            is_selected = selected_chart == chart_key
            btn_type = "primary" if is_selected else "secondary"
            if st.button(
                f"{chart_info['icon']}",
                key=f"chart_btn_{msg_id}_{chart_key}",
                help=f"{chart_info['name']}: {chart_info['description']}",
                type=btn_type,
                use_container_width=True
            ):
                st.session_state[f"chart_type_{msg_id}"] = chart_key
                st.rerun()
    
    st.caption(f"Selected: **{CHART_TYPES[selected_chart]['name']}** - {CHART_TYPES[selected_chart]['description']}")
    
    st.markdown("---")
    
    all_cols = df.columns.tolist()
    numeric_cols = viz_info["numeric_cols"]
    text_cols = viz_info["text_cols"]
    
    col1, col2, col3 = st.columns(3)
    
    with col1:
        x_options = text_cols + numeric_cols if text_cols else all_cols
        x_col = st.selectbox(
            "📍 X-Axis (Categories)",
            options=x_options,
            key=f"x_col_{msg_id}",
            help="Usually categories, dates, or labels"
        )
    
    with col2:
        y_options = numeric_cols if numeric_cols else all_cols
        y_col = st.selectbox(
            "📊 Y-Axis (Values)",
            options=y_options,
            key=f"y_col_{msg_id}",
            help="Usually numbers to measure"
        )
    
    with col3:
        color_options = ["None"] + text_cols
        color_col = st.selectbox(
            "🎨 Color By (Optional)",
            options=color_options,
            key=f"color_col_{msg_id}",
            help="Add color grouping"
        )
    
    if st.button("✨ Generate Chart", key=f"gen_btn_{msg_id}", type="primary", use_container_width=True):
        chart_type = st.session_state.get(f"chart_type_{msg_id}", "bar")
        color = color_col if color_col != "None" else None
        
        fig = generate_chart(df, chart_type, x_col, y_col, color)
        st.session_state[f"generated_chart_{msg_id}"] = fig
        st.rerun()
    
    if f"generated_chart_{msg_id}" in st.session_state:
        fig = st.session_state[f"generated_chart_{msg_id}"]
        st.plotly_chart(fig, use_container_width=True)
        
        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            html_buffer = io.StringIO()
            fig.write_html(html_buffer)
            st.download_button(
                "📥 Download Chart (HTML)",
                data=html_buffer.getvalue(),
                file_name=f"chart_{msg_id}.html",
                mime="text/html",
                key=f"dl_html_{msg_id}",
                use_container_width=True
            )
        with dl_col2:
            if st.button("🗑️ Clear Chart", key=f"clear_chart_{msg_id}", use_container_width=True):
                del st.session_state[f"generated_chart_{msg_id}"]
                st.rerun()
    
    return None

# ═══════════════════════════════════════════════════════
#  ✨ NEW: FUSION-SPECIFIC UI COMPONENTS
# ═══════════════════════════════════════════════════════

def render_guided_command_center():
    """Render the guided first screen for the command-center demo."""
    st.markdown(
        """
        <style>
            .fusion-command {
                border: 1px solid rgba(148, 163, 184, 0.26);
                border-radius: 8px;
                padding: 24px;
                background:
                    linear-gradient(135deg, rgba(15, 23, 42, 0.98), rgba(17, 24, 39, 0.96)),
                    radial-gradient(circle at 14% 0%, rgba(14, 165, 233, 0.22), transparent 34%);
                color: #f8fafc;
                margin-top: 12px;
            }
            .fusion-command h2 {
                margin: 0;
                color: #f8fafc;
                font-size: clamp(1.7rem, 3vw, 2.6rem);
                line-height: 1.08;
                letter-spacing: 0;
            }
            .fusion-command p {
                max-width: 820px;
                margin: 12px 0 0;
                color: #cbd5e1;
                line-height: 1.58;
                font-size: 1rem;
            }
            .fusion-chip-row {
                display: flex;
                flex-wrap: wrap;
                gap: 8px;
                margin-top: 16px;
            }
            .fusion-chip {
                border: 1px solid rgba(56, 189, 248, 0.28);
                border-radius: 999px;
                padding: 6px 10px;
                color: #e0f2fe;
                background: rgba(8, 47, 73, 0.42);
                font-size: 0.78rem;
                font-weight: 700;
            }
            .fusion-agent-flow {
                display: grid;
                grid-template-columns: repeat(5, minmax(116px, 1fr));
                gap: 10px;
                margin: 16px 0 4px;
            }
            .fusion-agent-step {
                min-height: 104px;
                padding: 13px;
                border-radius: 8px;
                border: 1px solid rgba(14, 116, 144, 0.22);
                background: #f8fafc;
            }
            .fusion-agent-step b {
                display: block;
                color: #0f172a;
                font-size: 0.9rem;
                margin-bottom: 5px;
            }
            .fusion-agent-step small {
                color: #64748b;
                line-height: 1.35;
            }
            .fusion-preview {
                border-left: 5px solid #0ea5e9;
                border-radius: 8px;
                padding: 18px 20px;
                background: #f8fafc;
                margin-top: 10px;
            }
            .fusion-preview h4 {
                margin: 0 0 8px;
                color: #0f172a;
            }
            .fusion-preview p {
                margin: 6px 0;
                color: #334155;
                line-height: 1.48;
            }
            .fusion-prompt-card {
                border: 1px solid rgba(148, 163, 184, 0.22);
                border-radius: 8px;
                padding: 15px;
                min-height: 126px;
                background: #f8fafc;
            }
            .fusion-prompt-card strong {
                display: block;
                color: #0f172a;
                margin-bottom: 5px;
            }
            .fusion-prompt-card p {
                margin: 0;
                color: #475569;
                font-size: 0.86rem;
                line-height: 1.42;
            }
            @media (max-width: 900px) {
                .fusion-command {
                    padding: 20px;
                }
                .fusion-agent-flow {
                    grid-template-columns: 1fr;
                }
            }
        </style>
        """,
        unsafe_allow_html=True,
    )

    st.markdown(
        """
        <div class="fusion-command">
            <h2>Guided Intelligence Command Center</h2>
            <p>
                Start with a guided question that proves the system is more than a chatbot:
                it routes across structured data, business documents, and web sources, then
                explains how trustworthy the answer is.
            </p>
            <div class="fusion-chip-row">
                <div class="fusion-chip">90,500 transactions</div>
                <div class="fusion-chip">23 business PDFs</div>
                <div class="fusion-chip">Hybrid BM25 + vector RAG</div>
                <div class="fusion-chip">SQL + RAG validation</div>
                <div class="fusion-chip">Live web pricing</div>
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    flow_col, preview_col = st.columns([1.1, 0.9])
    with flow_col:
        st.markdown(
            """
            <div class="fusion-agent-flow">
                <div class="fusion-agent-step"><b>SQL Agent</b><small>Queries exact transaction facts.</small></div>
                <div class="fusion-agent-step"><b>RAG Agent</b><small>Retrieves from indexed business docs.</small></div>
                <div class="fusion-agent-step"><b>Web Agent</b><small>Checks live competitor context.</small></div>
                <div class="fusion-agent-step"><b>Fusion Agent</b><small>Reconciles outputs and conflicts.</small></div>
                <div class="fusion-agent-step"><b>Validated Answer</b><small>Shows confidence, route, and sources.</small></div>
            </div>
            """,
            unsafe_allow_html=True,
        )
    with preview_col:
        st.markdown(
            """
            <div class="fusion-preview">
                <h4>Featured answer preview</h4>
                <p><b>Question:</b> Validate Q4 Electronics revenue across SQL and PDF reports.</p>
                <p><b>Expected result:</b> SQL and RAG validate about <b>$15.2M</b>, then show source difference and confidence.</p>
                <p><b>Evidence:</b> sales_transactions + Q4 2024 Financial Report.</p>
            </div>
            """,
            unsafe_allow_html=True,
        )

    st.markdown("### Run a Guided Demo")
    prompt_cols = st.columns(3)
    for idx, prompt in enumerate(RECRUITER_DEMO_PROMPTS):
        with prompt_cols[idx % 3]:
            st.markdown(
                f"""
                <div class="fusion-prompt-card">
                    <strong>{prompt["title"]}</strong>
                    <p>{prompt["proof"]}<br><b>Route:</b> {prompt["route"]}</p>
                </div>
                """,
                unsafe_allow_html=True,
            )
            if st.button("Run prompt", key=f"command_center_prompt_{idx}", use_container_width=True):
                st.session_state.pending_suggestion = prompt["question"]
                st.session_state.show_fusion_command_center = False
                st.rerun()

    st.info("Use Auto routing for the best demo path. Manual SQL, RAG, or Web modes are available in the sidebar for deeper inspection.")

def render_routing_badge(source_type: str):
    """
    Display which source(s) the Fusion Agent used
    
    Args:
        source_type: "sql_only" | "rag_only" | "web_only" | "sql_rag" | "comparison" etc.
    """
    
    route_config = {
        "no_data": {"icon": "🚫", "label": "No Data Available", "color": "#F44336"},
        "sql_only": {"icon": "🗄️", "label": "SQL Database", "color": "#4CAF50"},
        "rag_only": {"icon": "📄", "label": "PDF Documents", "color": "#2196F3"},
        "web_only": {"icon": "🌐", "label": "Web Scraping", "color": "#FF9800"},
        "comparison": {"icon": "🧠", "label": "RAG Comparison Mode", "color": "#9C27B0"},
        "sql_rag": {"icon": "🔗", "label": "SQL + RAG Fusion", "color": "#00BCD4"},
        "sql_web": {"icon": "🔗", "label": "SQL + Web Fusion", "color": "#FF5722"},
        "rag_web": {"icon": "🔗", "label": "RAG + Web Fusion", "color": "#795548"},
        "all": {"icon": "🌟", "label": "All Sources Fusion", "color": "#E91E63"},
    }
    
    config = route_config.get(source_type, {"icon": "❓", "label": source_type, "color": "#9E9E9E"})
    
    st.markdown(
        f"""
        <div style='
            background: linear-gradient(135deg, {config['color']}22 0%, {config['color']}11 100%);
            border-left: 4px solid {config['color']};
            padding: 12px 16px;
            border-radius: 8px;
            margin: 10px 0;
        '>
            <div style='display: flex; align-items: center; gap: 10px;'>
                <span style='font-size: 24px;'>{config['icon']}</span>
                <div>
                    <div style='font-weight: 600; color: #333; font-size: 14px;'>
                        ROUTING DECISION
                    </div>
                    <div style='color: {config['color']}; font-weight: 700; font-size: 16px; margin-top: 2px;'>
                        {config['label']}
                    </div>
                </div>
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )

def render_confidence_badge(validation: dict):
    """
    Display cross-validation confidence badge
    
    Args:
        validation: dict with 'confidence', 'confidence_reason', 'matches', 'discrepancies'
    """
    
    if not validation:
        return
    
    confidence = validation.get('confidence', 'UNKNOWN')
    reason = validation.get('confidence_reason', '')
    matches = validation.get('matches', [])
    discrepancies = validation.get('discrepancies', [])
    
    # Confidence colors
    confidence_config = {
        "HIGH": {"emoji": "✅", "color": "#4CAF50", "bg": "#E8F5E9"},
        "MEDIUM": {"emoji": "🟡", "color": "#FF9800", "bg": "#FFF3E0"},
        "LOW": {"emoji": "🔴", "color": "#F44336", "bg": "#FFEBEE"},
    }
    
    config = confidence_config.get(confidence, {"emoji": "⚪", "color": "#9E9E9E", "bg": "#F5F5F5"})
    
    st.markdown(
        f"""
        <div style='
            background: {config['bg']};
            border: 2px solid {config['color']};
            padding: 12px 16px;
            border-radius: 8px;
            margin: 10px 0;
        '>
            <div style='display: flex; align-items: center; gap: 10px; margin-bottom: 8px;'>
                <span style='font-size: 24px;'>{config['emoji']}</span>
                <div>
                    <div style='font-weight: 600; color: #666; font-size: 12px;'>
                        CROSS-VALIDATION CONFIDENCE
                    </div>
                    <div style='color: {config['color']}; font-weight: 700; font-size: 18px;'>
                        {confidence}
                    </div>
                </div>
            </div>
            <div style='color: #555; font-size: 14px; margin-top: 8px;'>
                {reason}
            </div>
        </div>
        """,
        unsafe_allow_html=True
    )
    
    # Show match/discrepancy details
    if matches or discrepancies:
        with st.expander("🔍 Validation Details", expanded=False):
            if matches:
                fact_label = "fact" if len(matches) == 1 else "facts"
                st.markdown(f"**✅ Validated {fact_label.title()} ({len(matches)}):**")
                for match in matches[:5]:  # Show top 5
                    sql_val = match.get('sql_value', 'N/A')
                    rag_val = match.get('rag_value', 'N/A')
                    pct_diff = match.get('pct_difference', 0)
                    st.markdown(f"- **{match.get('label', 'value')}**: SQL={sql_val:,.0f} ≈ RAG={rag_val:,.0f} (Δ{pct_diff:.2f}%)")
            
            if discrepancies:
                st.markdown(f"**⚠️ Discrepancies ({len(discrepancies)}):**")
                for disc in discrepancies[:3]:
                    st.markdown(f"- **{disc.get('label', 'value')}**: SQL={disc.get('sql_value', 'N/A')} vs RAG={disc.get('rag_value', 'N/A')}")

def render_sql_section(msg_id: str, sql_result: dict, is_latest: bool = False):
    """
    Render SQL results section (query + table + explanation + chart builder)
    
    Args:
        msg_id: Unique message ID
        sql_result: SQL agent output dict
        is_latest: Whether this is the latest message (auto-expand charts)
    """
    pd = _get_pd()
    
    if not sql_result or not sql_result.get('success'):
        st.warning("❌ SQL query failed or returned no results")
        if sql_result and sql_result.get('error'):
            st.error(f"Error: {sql_result['error']}")
        return
    
    with st.expander("🗄️ SQL Database Results", expanded=True):
        # SQL Query
        if sql_result.get('query'):
            st.markdown("**📝 Generated SQL Query:**")
            st.code(sql_result['query'], language="sql")
        
        # SQL Explanation
        if sql_result.get('explanation'):
            with st.expander("📖 How This Query Works"):
                st.markdown(sql_result['explanation'])
        
        # Data Table + Exports
        if sql_result.get('results'):
            df = pd.DataFrame(sql_result['results'])
            
            row_count = sql_result.get('row_count', len(df))
            row_label = "result row" if row_count == 1 else "result rows"
            st.markdown(f"**📊 Data Table ({row_count} {row_label}):**")
            st.dataframe(df, use_container_width=True)
            render_kpi_summary(df)
            
            # Export buttons
            st.markdown("**📥 Export Data:**")
            e1, e2, e3, e4 = st.columns(4)
            
            with e1:
                st.download_button(
                    "📄 CSV", data=df.to_csv(index=False),
                    file_name=f"fusion_sql_{msg_id}.csv", mime="text/csv",
                    use_container_width=True, key=f"sql_csv_{msg_id}"
                )
            
            with e2:
                st.download_button(
                    "📋 JSON",
                    data=df.to_json(orient="records", indent=2),
                    file_name=f"fusion_sql_{msg_id}.json",
                    mime="application/json",
                    use_container_width=True, key=f"sql_json_{msg_id}"
                )
            
            with e3:
                from openpyxl.utils import get_column_letter
                buf = io.BytesIO()
                with pd.ExcelWriter(buf, engine='openpyxl') as writer:
                    df.to_excel(writer, index=False, sheet_name='Results')
                    ws = writer.sheets['Results']
                    for ci, col_name in enumerate(df.columns, 1):
                        ml = max(
                            len(str(col_name)),
                            max((len(str(v)) for v in df[col_name]), default=0)
                        )
                        ws.column_dimensions[get_column_letter(ci)].width = min(ml + 3, 50)
                st.download_button(
                    "📊 Excel", data=buf.getvalue(),
                    file_name=f"fusion_sql_{msg_id}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True, key=f"sql_excel_{msg_id}"
                )
            
            with e4:
                st.download_button(
                    "📝 MD", data=df.to_markdown(index=False),
                    file_name=f"fusion_sql_{msg_id}.md", mime="text/markdown",
                    use_container_width=True, key=f"sql_md_{msg_id}"
                )
            
            # Chart Builder
            viz_info = can_visualize(df)
            
            if viz_info["can_chart"]:
                st.markdown("---")
                with st.expander("📊 Visualize This Data?", expanded=is_latest):
                    render_chart_builder(f"sql_{msg_id}", df)
            else:
                st.caption(f"📊 {viz_info['reason']}")
        
        # Timing info
        if sql_result.get('time'):
            st.caption(f"⏱️ SQL execution: {format_time(sql_result['time'])}")

def render_rag_section(msg_id: str, rag_result: dict):
    """
    Render RAG results section (answer + sources + document references)
    
    Args:
        msg_id: Unique message ID
        rag_result: RAG agent output dict
    """
    
    if not rag_result or not rag_result.get('success'):
        st.warning("❌ RAG query failed or found no relevant documents")
        if rag_result and rag_result.get('error'):
            st.error(f"Error: {rag_result['error']}")
        return
    
    with st.expander("📄 Document Search Results", expanded=True):
        # RAG Answer
        st.markdown("**💬 Document-Based Answer:**")
        st.markdown(rag_result['answer'])
        
        # Source Citations
        if rag_result.get('sources'):
            st.markdown("---")
            st.markdown(f"**📚 Source Documents ({len(rag_result['sources'])} citations):**")
            
            for i, source in enumerate(rag_result['sources'], 1):
                cited = "✅" if source.get('cited_in_answer') else "📎"
                similarity = source.get('similarity', '')
                sim_text = f" (Score: {similarity})" if similarity else ""
                
                st.markdown(
                    f"{cited} **{source['filename']}** (Page {source['page']}){sim_text}"
                )
        
        # Metadata
        col1, col2, col3 = st.columns(3)
        with col1:
            st.metric("Chunks Retrieved", rag_result.get('chunks_retrieved', 0))
        with col2:
            query_type = rag_result.get('query_type', 'simple')
            st.metric("Query Type", query_type.title())
        with col3:
            st.metric("Time", f"{rag_result.get('time', 0):.2f}s")

def render_web_section(msg_id: str, web_result: dict):
    """
    Render Web scraping results section with scraper status dashboard.
    """

    if not web_result or not web_result.get('success'):
        st.warning("❌ Web scraping failed or returned no data")
        if web_result and web_result.get('error'):
            st.error(f"Error: {web_result['error']}")
        return

    with st.expander("🌐 Competitor Intelligence", expanded=True):

        # ── Scraper Status Dashboard ──────────────────────────────
        raw_data = web_result.get('raw_data', {})
        statuses = raw_data.get('scraper_statuses', [])

        if statuses:
            st.markdown("**🔍 Scraper Status Dashboard**")
            cols = st.columns(len(statuses))
            for col, s in zip(cols, statuses):
                name    = s.get('name', 'Unknown')
                status  = s.get('status', 'unknown')
                products = s.get('products', 0)
                elapsed  = s.get('time', 0)
                error    = s.get('error')

                if status == 'success':
                    icon = '🟢'
                    label = f"{products} products"
                elif status == 'fallback':
                    icon = '🟡'
                    label = f"Mock ({products} items)"
                elif status == 'empty':
                    icon = '🟠'
                    label = 'No products'
                else:
                    icon = '🔴'
                    label = 'Failed'

                with col:
                    st.metric(
                        label=f"{icon} {name}",
                        value=label,
                        delta=f"{elapsed}s" if elapsed else None,
                        delta_color="off"
                    )
                    if error and status != 'success':
                        st.caption(f"_{error[:60]}_")

            # Show fallback warning if any mock data was used
            fallbacks = [s for s in statuses if s.get('status') == 'fallback']
            if fallbacks:
                st.warning("⚠️ All live scrapers failed — answer is based on sample data, not real prices.")

            st.markdown("---")

        # ── Competitor Analysis Answer ────────────────────────────
        st.markdown("**🛒 Competitor Analysis:**")
        if web_result.get('llm_error'):
            st.caption(f"⚠️ LLM unavailable — showing raw scraped prices. ({web_result['llm_error'][:80]})")
        st.markdown(web_result['answer'])

        # ── Product Details per Competitor ────────────────────────
        if raw_data.get('competitors'):
            st.markdown("---")
            st.markdown("**📊 Scraped Products:**")
            for comp_data in raw_data['competitors']:
                competitor = comp_data.get('competitor', 'Unknown')
                method     = comp_data.get('method', 'Unknown')
                total      = comp_data.get('total_found', len(comp_data.get('products', [])))
                products   = comp_data.get('products', [])
                is_mock    = comp_data.get('is_mock', False)

                badge = " *(sample data)*" if is_mock else ""
                st.markdown(f"**{competitor}**{badge} — {method}")
                st.caption(f"Found: {total} | Showing: {len(products)}")

                for i, product in enumerate(products[:5], 1):
                    st.markdown(f"{i}. **{product.get('name', 'Unknown')}** — {product.get('price', 'N/A')}")

                st.markdown("---")

        # ── Export ────────────────────────────────────────────────
        if raw_data:
            st.download_button(
                "📋 Download JSON",
                data=json.dumps(raw_data, indent=2, default=str),
                file_name=f"fusion_web_{msg_id}.json",
                mime="application/json",
                key=f"web_json_{msg_id}"
            )

        if web_result.get('time'):
            st.caption(f"⏱️ Web scraping: {format_time(web_result['time'])}")

def render_model_journey(models_tried: list):
    """
    Display which models were tried and their status
    
    Args:
        models_tried: List of dicts with 'model', 'status', 'time', 'error'
    """
    
    if not models_tried:
        return
    
    with st.expander("🔄 Model Journey", expanded=False):
        st.markdown("**Models attempted for this query:**")
        
        for i, attempt in enumerate(models_tried, 1):
            model = attempt.get('model', 'Unknown')
            status = attempt.get('status', '❓ UNKNOWN')
            time_taken = attempt.get('time', 0)
            error = attempt.get('error', '')
            
            # Color code by status
            if '✅' in status:
                color = "#4CAF50"
            elif '⏭️' in status:
                color = "#9E9E9E"
            else:
                color = "#F44336"
            
            st.markdown(
                f"""
                <div style='
                    border-left: 4px solid {color};
                    padding: 8px 12px;
                    margin: 8px 0;
                    background: {color}11;
                    border-radius: 4px;
                '>
                    <div style='font-weight: 600;'>{status} {model}</div>
                    <div style='font-size: 12px; color: #666; margin-top: 4px;'>
                        Time: {time_taken:.2f}s
                        {f"<br>Error: {error[:100]}" if error else ""}
                    </div>
                </div>
                """,
                unsafe_allow_html=True
            )

def _load_trace(trace_path: str) -> dict:
    if not trace_path:
        return {}
    try:
        path = Path(trace_path)
        if not path.exists():
            return {}
        return json.loads(path.read_text())
    except Exception:
        return {}

def _format_span_duration(seconds) -> str:
    try:
        value = float(seconds or 0)
    except (TypeError, ValueError):
        value = 0
    return f"{value:.2f}s"

def render_observability_panel(msg: dict):
    """Render local trace metadata for a Fusion response."""
    trace_id = msg.get("trace_id")
    trace_path = msg.get("trace_path")
    if not trace_id and not trace_path:
        return

    trace = _load_trace(trace_path)
    final = trace.get("final") or {}
    spans = trace.get("spans") or []
    diagnostics = get_trace_diagnostics(trace, slow_threshold=SLOW_SPAN_SECONDS)
    slowest = diagnostics["slowest_span"]
    slow_spans = diagnostics["slow_spans"]
    error_spans = diagnostics["error_spans"]

    route = final.get("source_type") or msg.get("source_type", "unknown")
    total_duration = trace.get("duration_s") or msg.get("query_time", 0)
    routing_model = final.get("routing_model") or msg.get("routing_model") or "n/a"
    validation = final.get("validation") or msg.get("validation") or {}
    validation_label = validation.get("confidence") or "n/a"

    with st.expander("🧭 How NexusIQ Ran This Answer", expanded=False):
        metric_cols = st.columns(4)
        metric_cols[0].metric("Route", route)
        metric_cols[1].metric("Total Time", _format_span_duration(total_duration))
        metric_cols[2].metric("Validation", validation_label)
        metric_cols[3].metric("Router", str(routing_model))

        if slowest:
            slow_label = "⚠️ " if (slowest.get("duration_s") or 0) > SLOW_SPAN_SECONDS else ""
            st.caption(
                f"{slow_label}Slowest step: **{slowest.get('name')}** "
                f"({_format_span_duration(slowest.get('duration_s'))})"
            )

        if error_spans:
            st.warning(
                "One or more trace steps recorded an error: "
                + ", ".join(span.get("name", "unknown") for span in error_spans)
            )
        elif slow_spans:
            st.info(
                "Slow steps over 3s: "
                + ", ".join(
                    f"{span.get('name')} ({_format_span_duration(span.get('duration_s'))})"
                    for span in slow_spans
                )
            )

        if spans:
            st.markdown("**Trace timeline**")
            for span in spans:
                status_icon = "✅" if span.get("status") == "ok" else "⚠️"
                slow_icon = " ⏳" if (span.get("duration_s") or 0) > SLOW_SPAN_SECONDS else ""
                st.caption(
                    f"{status_icon}{slow_icon} {span.get('name')} · "
                    f"{_format_span_duration(span.get('duration_s'))}"
                )

        trace_col, path_col = st.columns([0.4, 0.6])
        trace_col.caption(f"Trace ID: `{trace_id or trace.get('trace_id')}`")
        if trace_path:
            path_col.caption(f"Trace file: `{trace_path}`")

# ─────────────────────────────────────────────────────
#  ✨ NEW: FUSION MESSAGE RENDERER
# ─────────────────────────────────────────────────────

def render_fusion_message(msg: dict, is_latest: bool = False):
    """
    Render a complete Fusion Agent message with all sections
    
    Args:
        msg: Message dict with 'answer', 'source_type', 'sql_result', 'rag_result', 'web_result', etc.
        is_latest: Whether this is the latest message
    """
    
    msg_id = msg.get("id", "0")
    
    # ═══════════════════════════════════════════════════════════
    # 1. ROUTING DECISION
    # ═══════════════════════════════════════════════════════════
    
    render_routing_badge(msg.get("source_type", "unknown"))

    # Routing fallback warning
    if msg.get("routing_fallback"):
        routing_model = msg.get("routing_model", "backup model")
        st.warning(
            f"⚠️ **Routing used fallback LLM** ({routing_model}) — primary model was unavailable. "
            "Routing decision may differ from normal. Results could vary.",
            icon="⚠️"
        )

    # ═══════════════════════════════════════════════════════════
    # 2. FUSED ANSWER (Main response)
    # ═══════════════════════════════════════════════════════════
    
    st.markdown("### 💡 Answer")
    st.markdown(msg.get("answer", "No answer available"))
    
    # ═══════════════════════════════════════════════════════════
    # 3. CROSS-VALIDATION (if available)
    # ═══════════════════════════════════════════════════════════
    
    if msg.get("validation"):
        render_confidence_badge(msg["validation"])
    
    st.markdown("---")
    
    # ═══════════════════════════════════════════════════════════
    # 4. SOURCE-SPECIFIC SECTIONS (Expandable)
    # ═══════════════════════════════════════════════════════════
    
    # SQL Section
    if msg.get("sql_result"):
        render_sql_section(msg_id, msg["sql_result"], is_latest)
    
    # RAG Section
    if msg.get("rag_result"):
        render_rag_section(msg_id, msg["rag_result"])
    
    # Web Section
    if msg.get("web_result"):
        render_web_section(msg_id, msg["web_result"])
    
    # ═══════════════════════════════════════════════════════════
    # 5. MODEL JOURNEY (All models tried across all agents)
    # ═══════════════════════════════════════════════════════════
    
    all_models_tried = []
    
    if msg.get("sql_result") and msg["sql_result"].get("models_tried"):
        all_models_tried.extend(msg["sql_result"]["models_tried"])
    
    if msg.get("rag_result") and msg["rag_result"].get("models_tried"):
        all_models_tried.extend(msg["rag_result"]["models_tried"])
    
    if all_models_tried:
        render_model_journey(all_models_tried)

    # ═══════════════════════════════════════════════════════════
    # 6. OBSERVABILITY TRACE
    # ═══════════════════════════════════════════════════════════

    render_observability_panel(msg)
    
    # ═══════════════════════════════════════════════════════════
    # 7. TIMING INFO
    # ═══════════════════════════════════════════════════════════
    
    total_time = msg.get("query_time", 0)
    from_cache = msg.get("from_cache", False)

    if from_cache:
        st.caption(f"⚡ Returned from cache in {total_time:.2f}s (original query was slower)")
    else:
        st.caption(f"⏱️ Total query time: {format_time(total_time)}")

# ═══════════════════════════════════════════════════════
#  INITIALIZE AGENT
# ═══════════════════════════════════════════════════════

@st.cache_resource(show_spinner=False)
def get_agent():
    from agents.fusion_agent import get_fusion_agent    # ✅ lazy import
    return get_fusion_agent()



def add_to_history(question, result, execution_time):
    """Add query to history (now stores full fusion result)"""
    st.session_state.query_history.insert(0, {
        "question": question,
        "answer": result.get("answer", ""),
        "source_type": result.get("source_type", "unknown"),
        "sql_result": result.get("sql_result"),
        "rag_result": result.get("rag_result"),
        "web_result": result.get("web_result"),
        "validation": result.get("validation"),
        "sources": result.get("sources", []),
        "trace_id": result.get("trace_id"),
        "trace_path": result.get("trace_path"),
        "routing_model": result.get("routing_model"),
        "routing_fallback": result.get("routing_fallback"),
        "time": execution_time,
        "timestamp": datetime.now(),
        "success": True
    })
    st.session_state.query_history = st.session_state.query_history[:10]

# ═══════════════════════════════════════════════════════
#  PAGE LAYOUT
# ═══════════════════════════════════════════════════════


def _scroll_to_anchor(anchor_id: str, offset: int = 16):
    import streamlit.components.v1 as components
    script = """
        <script>
            function scrollToAnchor() {
                var d = window.parent.document;
                var anchor = d.getElementById('__ANCHOR_ID__');
                if (!anchor) return false;

                var scrollers = Array.from(d.querySelectorAll('*')).filter(function(el) {
                    return el.scrollHeight > el.clientHeight + 20;
                });
                var main = scrollers.find(function(el) {
                    return el.matches && (
                        el.matches('section.stMain') ||
                        el.getAttribute('data-testid') === 'stAppScrollToBottomContainer'
                    );
                }) || d.scrollingElement || d.documentElement;

                var mainRect = main.getBoundingClientRect ? main.getBoundingClientRect() : {top: 0};
                var anchorRect = anchor.getBoundingClientRect();
                var currentTop = main === d.scrollingElement || main === d.documentElement
                    ? (window.parent.pageYOffset || d.documentElement.scrollTop || 0)
                    : main.scrollTop;
                var targetTop = currentTop + anchorRect.top - mainRect.top - __OFFSET__;

                if (main.scrollTo) {
                    main.scrollTo({top: Math.max(0, targetTop), behavior: 'smooth'});
                } else {
                    anchor.scrollIntoView({behavior: 'smooth', block: 'start'});
                }
                return true;
            }
            scrollToAnchor();
            setTimeout(scrollToAnchor, 150);
            setTimeout(scrollToAnchor, 350);
            setTimeout(scrollToAnchor, 700);
        </script>
    """
    components.html(
        script.replace("__ANCHOR_ID__", anchor_id).replace("__OFFSET__", str(offset)),
        height=1,
        scrolling=False,
    )

def _scroll_to_bottom():
    _scroll_to_anchor("nexusiq-latest-answer", offset=20)

def _scroll_to_question():
    _scroll_to_anchor("nexusiq-latest-question", offset=20)

def _scroll_to_command_center():
    _scroll_to_anchor("nexusiq-command-center", offset=20)

def _scroll_to_top():
    import streamlit.components.v1 as components
    components.html("""
        <script>
            function scrollToTop() {
                var d = window.parent.document;
                var main = d.querySelector('section.stMain') || d.querySelector('[data-testid="stAppViewContainer"]');
                if (main) {
                    main.scrollTo({top: 0, behavior: 'smooth'});
                    return true;
                }
                window.parent.scrollTo({top: 0, behavior: 'smooth'});
                return false;
            }
            scrollToTop();
            setTimeout(scrollToTop, 200);
            setTimeout(scrollToTop, 500);
        </script>
    """, height=1, scrolling=False)

def run_fusion_chat():
    """Main function for Fusion Chat interface"""
    from streamlit.runtime.scriptrunner import add_script_run_ctx, get_script_run_ctx

    # ═══════════════════════════════════════════════════════
    #  LOAD AGENT (cached — only slow on first load)
    # ═══════════════════════════════════════════════════════

    if "nexusiq_agent" not in st.session_state:
        # Kick off the heavy load on a background thread so the script
        # thread can stay in a progress-update loop. `time.sleep` inside
        # the loop releases the GIL, which is what lets Streamlit's
        # message pump actually push deltas to the browser during a
        # long-running init — that's the canonical Streamlit progress
        # pattern.
        if "_agent_loader" not in st.session_state:
            _ctx = get_script_run_ctx()
            _result: dict = {}

            def _worker():
                add_script_run_ctx(threading.current_thread(), _ctx)
                try:
                    _result["agent"] = get_agent()
                except Exception as _exc:
                    _result["error"] = _exc

            _t = threading.Thread(target=_worker, daemon=True)
            _t.start()
            st.session_state._agent_loader = (_t, _result)

        _thread, _result = st.session_state._agent_loader

        # Render loading UI directly (no placeholder wrapper — wrappers batch).
        st.markdown("<br><br>", unsafe_allow_html=True)
        _l, _c, _r = st.columns([1, 2, 1])
        with _c:
            st.markdown(
                """
                <div style='text-align:center; padding:40px;'>
                    <div style='font-size:72px; margin-bottom:16px;'>🧠</div>
                    <h2 style='color:#4F8BF9; margin-bottom:8px;'>Loading Fusion Agent</h2>
                    <p style='color:#888; font-size:16px;'>Initializing AI models & vector database...</p>
                    <p style='color:#aaa; font-size:13px; margin-top:8px;'>First load only — ~20 seconds</p>
                </div>
                """,
                unsafe_allow_html=True,
            )
            _progress = st.progress(0, text="Loading AI models…")

        _start = time.time()
        while _thread.is_alive():
            _elapsed = time.time() - _start
            _pct = min(int(_elapsed * 4), 95)
            _progress.progress(_pct, text=f"Loading AI models… ({int(_elapsed)}s)")
            time.sleep(0.2)

        _progress.progress(100, text="Ready!")

        if "error" in _result:
            st.error(f"Failed to load Fusion Agent: {_result['error']}")
            del st.session_state._agent_loader
            st.stop()

        st.session_state.nexusiq_agent = _result["agent"]
        del st.session_state._agent_loader
        st.rerun()

    agent = st.session_state.nexusiq_agent

    # ═══════════════════════════════════════════════════════
    #  SESSION STATE
    # ═══════════════════════════════════════════════════════

    if "query_history" not in st.session_state:
        st.session_state.query_history = []
    if "chat_messages" not in st.session_state:
        st.session_state.chat_messages = []
    if "pending_suggestion" not in st.session_state:
        st.session_state.pending_suggestion = None
    if "from_history" not in st.session_state:
        st.session_state.from_history = False
    if "pending_correction" not in st.session_state:
        st.session_state.pending_correction = None   # {"original": str, "corrected": str, "corrections": list}
    if "_last_corrected_q" not in st.session_state:
        st.session_state["_last_corrected_q"] = None
    if "source_filter" not in st.session_state:
        st.session_state.source_filter = "Auto"  # ✨ NEW: Default to auto-routing
    if "web_category" not in st.session_state:
        st.session_state.web_category = "electronics"  # ✨ NEW: Default web category
    if "show_fusion_command_center" not in st.session_state:
        st.session_state.show_fusion_command_center = False
    if "pending_query_to_process" not in st.session_state:
        st.session_state.pending_query_to_process = None
    if "scroll_target" not in st.session_state:
        st.session_state.scroll_target = None
    
    st.title("🔗 Fusion Agent — Multi-Source Intelligence")
    st.markdown("*Cross-validates answers across SQL database, business PDFs, and live competitor pricing*")
    st.markdown(
        "<p style='font-size:13px; color:#6b7280; margin-top:-8px;'>"
        "<code>Gemini 2.5</code> &nbsp;·&nbsp; <code>Groq LLaMA 3.3</code> &nbsp;·&nbsp; "
        "<code>PostgreSQL</code> &nbsp;·&nbsp; <code>ChromaDB</code> &nbsp;·&nbsp; "
        "<code>BM25 + Vector Search</code> &nbsp;·&nbsp; <code>Live Web Scraping</code>"
        "</p>",
        unsafe_allow_html=True,
    )
    st.caption(
        f"Routing: {st.session_state.source_filter} · "
        f"Web category: {st.session_state.web_category}"
    )
    
    # ═══════════════════════════════════════════════════════
    #  SIDEBAR
    # ═══════════════════════════════════════════════════════
    
    with st.sidebar:
        if st.button("🧭 Show Demo Guide", key="show_demo_guide_top", use_container_width=True):
            st.session_state.show_fusion_command_center = True
            st.session_state.scroll_target = None
            st.session_state.pending_query_to_process = None
            st.rerun()

        st.markdown("---")
        st.header("⚙️ Source Controls")
        
        # ✨ NEW: Source Filter
        st.subheader("🎯 Routing Mode")
        source_filter = st.radio(
            "Choose how to route queries:",
            ["Auto", "SQL Only", "RAG Only", "Web Only"],
            index=0,
            key="source_filter_radio",
            help="Auto: Let Fusion Agent decide | Manual: Force a specific source"
        )
        st.session_state.source_filter = source_filter
        
        # ✨ NEW: Web Category Selector (only show if Web is involved)
        if source_filter in ["Auto", "Web Only"]:
            st.markdown("---")
            st.subheader("🛒 Web Scraping Category")
            web_category = st.selectbox(
                "Product category for competitor data:",
                ["electronics", "home", "clothing", "food", "sports"],
                index=0,
                key="web_cat_select"
            )
            st.session_state.web_category = web_category
        
        with st.expander("📊 Database Schema", expanded=False):
            st.markdown("**📋 sales_transactions**")
            st.code(
                "• transaction_date\n• region (5 regions)\n• store_id\n"
                "• product_category\n• product_name\n• quantity, unit_price\n"
                "• total_amount\n• customer_id\n• payment_method"
            )
            st.markdown("**👥 customers**")
            st.code(
                "• customer_id\n• name, email, region\n"
                "• signup_date\n• total_purchases"
            )
        
        st.markdown("---")
        st.subheader("💡 Example Questions")
        
        example_questions = [
            ("What was Q4 2024 revenue?", "Fusion (SQL+RAG)"),
            ("Compare Q3 and Q4 performance", "RAG Comparison"),
            ("What are competitor prices for electronics?", "Web Scraping"),
            ("Top 5 products by revenue", "SQL Only"),
            ("What is the return policy?", "RAG Only"),
        ]
        
        for eq, hint in example_questions:
            if st.button(f"💬 {eq}", key=f"ex_{eq[:20]}", use_container_width=True):
                st.session_state.pending_suggestion = eq
                st.session_state.show_fusion_command_center = False
                st.rerun()
            st.caption(f"→ {hint}")
        
        with st.expander("📊 Model Status", expanded=False):
            # Show Gemini Pro status
            settings = _get_settings()
            if settings.use_gemini_pro:
                st.warning("🟡 Gemini Pro: **ENABLED** (may exhaust quickly)")
            else:
                st.info("🔵 Gemini Pro: **DISABLED** (free tier protection)")
            
            # Get quota status from SQL agent (Fusion uses same models)
            quota_status = agent.sql_agent.get_quota_status()
            if quota_status:
                for model, status in quota_status.items():
                    if "pro" in model.lower() and not settings.use_gemini_pro:
                        continue
                    st.caption(f"{status['status']} {model.split('-')[0]}")
            else:
                st.caption("🟢 All models available")
            
            if st.button("🔄 Reset Quota Tracking", use_container_width=True):
                agent.sql_agent.reset_quota_tracking()
                st.rerun()
        
        st.markdown("---")
        st.subheader("📜 Query History")
        
        if st.session_state.query_history:
            for i, item in enumerate(st.session_state.query_history[:5]):
                # Show source type icon
                source_type = item.get("source_type", "unknown")
                icon = SOURCE_ICONS.get(source_type.split('_')[0], "❓")
                
                short = item["question"][:25] + ("..." if len(item["question"]) > 25 else "")
                
                if st.button(f"{icon} {short}", key=f"hist_{i}", use_container_width=True):
                    st.session_state.pending_suggestion = item["question"]
                    st.session_state.show_fusion_command_center = False
                    st.session_state.from_history = True   # ← flag it
                    st.rerun()
                
                st.caption(f"⏱️ {item['time']:.1f}s • {time_ago(item['timestamp'])}")
            
            if st.button("🗑️ Clear All", use_container_width=True):
                st.session_state.query_history = []
                st.session_state.chat_messages = []
                st.session_state.show_fusion_command_center = False
                st.session_state.scroll_target = None
                # Clear chart state
                keys_to_clear = [k for k in st.session_state.keys() 
                                if k.startswith(('chart_', 'generated_chart_', 'x_col_', 'y_col_', 'color_col_'))]
                for k in keys_to_clear:
                    del st.session_state[k]
                st.rerun()
        else:
            st.caption("No queries yet")
    
    # ═══════════════════════════════════════════════════════
    #  REPLAY CHAT HISTORY
    # ═══════════════════════════════════════════════════════

    if st.session_state.get("show_fusion_command_center"):
        st.markdown("---")
        st.session_state.scroll_target = None
        st.markdown('<div id="nexusiq-command-center"></div>', unsafe_allow_html=True)
        _scroll_to_command_center()
        render_guided_command_center()
        st.stop()
    
    total_messages = len(st.session_state.chat_messages)
    
    for idx, msg in enumerate(st.session_state.chat_messages):
        is_latest = (idx == total_messages - 1) and msg["role"] == "assistant"
        is_latest_user = (idx == total_messages - 1) and msg["role"] == "user"
        
        if msg["role"] == "user":
            if is_latest_user:
                st.markdown('<div id="nexusiq-latest-question"></div>', unsafe_allow_html=True)
            with st.chat_message("user"):
                st.markdown(msg["content"])
        else:
            if is_latest:
                st.markdown('<div id="nexusiq-latest-answer"></div>', unsafe_allow_html=True)
            with st.chat_message("assistant", avatar="🔗"):
                render_fusion_message(msg, is_latest=is_latest)

    # Auto-scroll to latest message after history replay
    if st.session_state.get("scroll_target") == "question":
        _scroll_to_question()
        if st.session_state.get("pending_query_to_process"):
            time.sleep(0.7)
            st.session_state.scroll_target = None
            st.rerun()
        st.session_state.scroll_target = None
    elif st.session_state.get("scroll_target") == "answer":
        _scroll_to_bottom()
        st.session_state.scroll_target = None

    # ═══════════════════════════════════════════════════════
    #  DID YOU MEAN? (spell-correction prompt)
    # ═══════════════════════════════════════════════════════

    if st.session_state.pending_correction:
        corr = st.session_state.pending_correction
        correction_labels = " | ".join(
            [f"**{c['from']}** → **{c['to']}**" for c in corr["corrections"]]
        )
        st.markdown('<div id="nexusiq-latest-answer"></div>', unsafe_allow_html=True)
        _scroll_to_bottom()
        with st.chat_message("assistant", avatar="🔗"):
            st.info(
                f"Did you mean: **\"{corr['corrected']}\"**?\n\n"
                f"*Auto-detected correction: {correction_labels}*"
            )
            col1, col2 = st.columns(2)
            with col1:
                if st.button(
                    f"✅ Yes — run \"{corr['corrected']}\"",
                    key="use_corrected",
                    use_container_width=True
                ):
                    st.session_state.pending_suggestion = corr["corrected"]
                    st.session_state.pending_correction = None
                    st.rerun()
            with col2:
                if st.button(
                    f"❌ No — run as typed",
                    key="use_original",
                    use_container_width=True
                ):
                    st.session_state.pending_suggestion = corr["original"]
                    st.session_state.pending_correction = None
                    st.session_state["_last_corrected_q"] = corr["original"]
                    st.rerun()

    # ═══════════════════════════════════════════════════════
    #  HANDLE INPUT
    # ═══════════════════════════════════════════════════════

    if st.session_state.pending_query_to_process:
        question = st.session_state.pending_query_to_process
        st.session_state.pending_query_to_process = None
    elif st.session_state.pending_suggestion:
        question = st.session_state.pending_suggestion
        st.session_state.pending_suggestion = None
    else:
        question = st.chat_input("💬 Ask a question across all data sources...")
    
    # ═══════════════════════════════════════════════════════
    #  PROCESS NEW QUESTION
    # ═══════════════════════════════════════════════════════
    
    if question:
        from utils.validators import auto_correct_question
        st.session_state.show_fusion_command_center = False

        # Check for spelling corrections BEFORE running the query.
        # Only intercept on fresh user input — not when re-running a corrected query.
        is_corrected_rerun = question == (st.session_state.get("_last_corrected_q"))
        correction = auto_correct_question(question) if not is_corrected_rerun else {"corrected": False}

        if correction["corrected"] and not is_corrected_rerun:
            # Show the user's original message in chat
            st.session_state.chat_messages.append({"role": "user", "content": question})
            st.session_state.scroll_target = "question"
            with st.chat_message("user"):
                st.markdown(question)

            # Store correction for the "Did you mean?" UI and rerun
            st.session_state.pending_correction = {
                "original": question,
                "corrected": correction["corrected_question"],
                "corrections": correction["corrections"],
            }
            st.session_state["_last_corrected_q"] = correction["corrected_question"]
            st.rerun()

        # Mark this as a resolved corrected query so we don't re-intercept it
        if is_corrected_rerun:
            st.session_state["_last_corrected_q"] = None

        # Add user message
        already_visible = (
            st.session_state.chat_messages
            and st.session_state.chat_messages[-1].get("role") == "user"
            and st.session_state.chat_messages[-1].get("content") == question
        )
        if not already_visible:
            st.session_state.chat_messages.append({
                "role": "user",
                "content": question
            })
            st.session_state.pending_query_to_process = question
            st.session_state.scroll_target = "question"
            st.rerun()

        # Process with Fusion Agent
        with st.chat_message("assistant", avatar="🔗"):
            status = st.empty()
            insight_box = st.empty()
            sql_progress = st.empty()
            rag_progress = st.empty()
            web_progress = st.empty()

            insight_box.info(random.choice(INSIGHTS))
            status.markdown("🔍 Analyzing question and routing to sources...")

            source_filter = st.session_state.get("source_filter", "Auto")
            force_source_map = {
                "SQL Only": "sql_only",
                "RAG Only": "rag_only",
                "Web Only": "web_only",
            }
            force_source = force_source_map.get(source_filter)

            # Progressive disclosure: update placeholders as each agent finishes
            _ctx = get_script_run_ctx()
            def _progress_cb(source_name: str, agent_result: dict):
                from streamlit.runtime.scriptrunner import add_script_run_ctx
                add_script_run_ctx(threading.current_thread(), _ctx)
                ok = agent_result.get("success", False)
                t = agent_result.get("time", 0)
                icon = "✅" if ok else "⚠️"
                if source_name == "sql":
                    sql_progress.markdown(f"{icon} **SQL** — {t:.1f}s")
                elif source_name == "rag":
                    rag_progress.markdown(f"{icon} **RAG Docs** — {t:.1f}s")
                elif source_name == "web":
                    web_progress.markdown(f"{icon} **Web** — {t:.1f}s")

            start_time = time.time()
            result = agent.query(question, force_source=force_source, progress_cb=_progress_cb)
            total_time = time.time() - start_time

            # If result came from cache, override with actual retrieval time
            if result.get('_from_cache'):
                total_time = time.time() - start_time  # will be <1s
                result = {**result, 'query_time': total_time}  # override stored time

            status.empty()
            insight_box.empty()
            sql_progress.empty()
            rag_progress.empty()
            web_progress.empty()
            
            msg_id = str(int(time.time() * 1000))
            
            # Render the fusion result
            st.markdown('<div id="nexusiq-latest-answer"></div>', unsafe_allow_html=True)
            render_fusion_message({
                "id": msg_id,
                "answer": result.get("answer", "No answer generated"),
                "source_type": result.get("source_type", "unknown"),
                "sql_result": result.get("sql_result"),
                "rag_result": result.get("rag_result"),
                "web_result": result.get("web_result"),
                "validation": result.get("validation"),
                "sources": result.get("sources", []),
                "query_time": total_time,
                "from_cache": result.get('_from_cache', False),
                "trace_id": result.get("trace_id"),
                "trace_path": result.get("trace_path"),
                "routing_model": result.get("routing_model"),
                "routing_fallback": result.get("routing_fallback")
            }, is_latest=True)
            
            # Save to chat history
            st.session_state.chat_messages.append({
                "role": "assistant",
                "id": msg_id,
                "answer": result.get("answer", ""),
                "source_type": result.get("source_type", "unknown"),
                "sql_result": result.get("sql_result"),
                "rag_result": result.get("rag_result"),
                "web_result": result.get("web_result"),
                "validation": result.get("validation"),
                "sources": result.get("sources", []),
                "query_time": total_time,
                "trace_id": result.get("trace_id"),
                "trace_path": result.get("trace_path"),
                "routing_model": result.get("routing_model"),
                "routing_fallback": result.get("routing_fallback")
            })
            st.session_state.scroll_target = "answer"
            
            # Save to query history
            if not st.session_state.get("from_history", False):
                add_to_history(question, result, total_time)
        
        st.rerun()
    
    # ═══════════════════════════════════════════════════════
    #  EMPTY STATE (Welcome Screen)
    # ════════════════════════════════════
    if not st.session_state.chat_messages and not st.session_state.get("show_fusion_command_center"):
        st.markdown("---")
        st.markdown('<div id="nexusiq-command-center"></div>', unsafe_allow_html=True)
        _scroll_to_command_center()
        render_guided_command_center()

# Run the app
if __name__ == "__main__":
    run_fusion_chat()
